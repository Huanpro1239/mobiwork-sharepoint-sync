from __future__ import annotations

import json
import tempfile
from datetime import date, datetime
from pathlib import Path
from zoneinfo import ZoneInfo
from math import ceil

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from excel_export import write_workbook

LAYOUT = Path(__file__).resolve().parents[1] / "config/promotion_detail_layout.json"


def _format_tracking(sheet, month: date) -> None:
    from promotion_tracking import TITLE

    sheet.insert_rows(1, 4)
    sheet["A1"], sheet["C1"] = "Thời gian in:", datetime.now(ZoneInfo("Asia/Ho_Chi_Minh")).replace(tzinfo=None)
    sheet["C1"].number_format = "dd/mm/yyyy hh:mm:ss"
    sheet["A2"], sheet["C2"] = "Tháng báo cáo:", month.strftime("%m/%Y")
    sheet["A3"] = TITLE
    sheet["A4"] = "THÔNG BÁO: Tất cả CT trong kỳ; mỗi khách/mức là 1 suất. Ngày đăng ký thiếu để trống."
    for address in ("A1:B1", "A2:B2", "A3:M3", "A4:M4"):
        sheet.merge_cells(address)
    sheet["N3"] = "Tích lũy: tổng từ đầu CT. Còn lại: kỳ xét thưởng; âm là phần chưa đạt, 0 là đủ mục tiêu tối thiểu."
    sheet.merge_cells(f"N3:{get_column_letter(sheet.max_column)}3")
    sheet["N4"] = "Số suất đạt theo doanh số/sản lượng, chưa xác nhận chi thưởng. Tháng ngoài CT hoặc chưa có nguồn để trống."
    sheet.merge_cells(f"N4:{get_column_letter(sheet.max_column)}4")
    widths = [5, 18, 23, 24, 24, 8, 16, 27, 24, 9, 38, 16, 44, 17, 18, 18, 10, 20, 20, 18]
    headers = {cell.column: str(cell.value or "") for cell in sheet[5]}
    for row in sheet.iter_rows():
        for cell in row:
            cell.font = Font(name="Times New Roman", size=10, bold=cell.row in {3, 5})
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if cell.row == 5:
                cell.fill = PatternFill("solid", fgColor="CCFF99")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            elif cell.row >= 6:
                label = headers[cell.column]
                if cell.column >= 16 and label not in {"Mã CT", "ĐVT mục tiêu", "Cách tính", "Trạng thái trả thưởng", "Thông tin nguồn"}:
                    integer = isinstance(cell.value, (int, float)) and float(cell.value).is_integer()
                    cell.number_format = "#,##0" if integer else "#,##0.000000"
                if label == "NGÀY ĐĂNG KÝ":
                    cell.number_format = "dd/mm/yyyy"
    for column, label in headers.items():
        width = widths[column - 1] if column <= len(widths) else 19
        if label in {"Mã CT", "Trạng thái trả thưởng", "Thông tin nguồn"}:
            width = 42
        sheet.column_dimensions[get_column_letter(column)].width = width
    for row in sheet.iter_rows(min_row=6):
        lines = max(max(1, ceil(len(str(c.value or "")) / max(sheet.column_dimensions[c.column_letter].width - 2, 1))) for c in row)
        sheet.row_dimensions[row[0].row].height = min(400, max(30, 15 * lines))
    sheet["A3"].font = Font(name="Times New Roman", size=14, bold=True)
    sheet.row_dimensions[3].height = 30
    sheet.row_dimensions[4].height = 30
    sheet.row_dimensions[5].height = 48
    sheet.freeze_panes = "H6"
    sheet.auto_filter.ref = f"A5:{get_column_letter(sheet.max_column)}{sheet.max_row}"
    sheet.print_title_rows = "1:5"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_area = f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.tabColor = "006B54"


def _format_reward_views(workbook) -> None:
    """Keep long DMS names and statuses readable in the supplemental reward views."""
    for name in ("TraThuong", "KhuyenMaiDonHang", "ChuongTrinh", "Thuong_theo_don", "DoanhSoChuaTinh"):
        if name not in workbook:
            continue
        sheet = workbook[name]
        headers = {cell.column: str(cell.value or "") for cell in sheet[1]}
        sheet.row_dimensions[1].height = 36
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="CCFF99")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for row in sheet.iter_rows(min_row=2):
            lines = 1
            for cell in row:
                label = headers[cell.column]
                if isinstance(cell.value, str):
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                    width = max((sheet.column_dimensions[cell.column_letter].width or 10) - 2, 1)
                    lines = max(lines, sum(max(ceil(len(part) / width), 1) for part in cell.value.split("\n")))
                if label.endswith("(đ)"):
                    cell.number_format = "#,##0"
                elif label.startswith("Số lượng"):
                    cell.number_format = "#,##0.######"
                elif label in {"Từ ngày", "Đến ngày", "Ngày Đơn hàng"}:
                    cell.number_format = "dd/mm/yyyy"
            sheet.row_dimensions[row[0].row].height = min(400, max(18, 15 * lines))


def write_detail_workbook(frames, filename: str, month: date, output_dir: Path = Path("output"),
                          title: str | None = None) -> Path:
    """Production workbook writer: apply the user's sanitized template without customer data."""
    layout = json.loads(LAYOUT.read_text(encoding="utf-8"))
    if len(frames["BaoCao"]) + 4 > 1_048_576:
        raise ValueError("Promotion detail exceeds Excel row limit with template header")
    output_dir.mkdir(parents=True, exist_ok=True)
    destination = output_dir / filename
    with tempfile.TemporaryDirectory(dir=output_dir) as staging:
        path = write_workbook(frames, filename, Path(staging))
        workbook = load_workbook(path)
        try:
            _format_reward_views(workbook)
            for name in ("BaoCao", "KhuyenMaiDonHang"):
                if name in workbook:
                    _format_detail_sheet(workbook[name], layout, month, title if name == "BaoCao" else None)
            if "TheoDoiTichLuy" in workbook:
                _format_tracking(workbook["TheoDoiTichLuy"], month)
                workbook.active = workbook.sheetnames.index("TheoDoiTichLuy")
            workbook.save(path)
        finally:
            workbook.close()
        path.replace(destination)
    return destination


def _format_detail_sheet(sheet, layout, month, title=None) -> None:
    sheet.insert_rows(1, 3)
    sheet["A1"] = "Thời gian in:"
    sheet["B1"] = datetime.now(ZoneInfo("Asia/Ho_Chi_Minh")).strftime("%d/%m/%Y %H:%M:%S")
    sheet["A2"] = "Tháng báo cáo:"
    sheet["B2"] = month.strftime("%Y-%m")
    sheet["A3"] = title or layout["title"]
    sheet["Y3"] = "(Đơn vị: KÉT/THÙNG/BÌNH; vật phẩm tặng: CÁI)"
    for merged in layout["merged"]:
        sheet.merge_cells(merged)
    headers, widths = list(layout["headers"]), list(layout["widths"])
    for column in range(len(headers) + 1, sheet.max_column + 1):  # report-specific extra columns
        headers.append(str(sheet.cell(4, column).value or ""))
        widths.append(22.0)
    last_column = get_column_letter(len(headers))
    money_columns = {i for i, label in enumerate(headers, 1) if label.endswith("(đ)")}
    for column, (label, width) in enumerate(zip(headers, widths, strict=True), 1):
        cell = sheet.cell(4, column, label)
        cell.font = Font(name="Times New Roman", size=10, bold=True)
        cell.fill = PatternFill("solid", fgColor="CCFF99")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet["A3"].font = Font(name="Times New Roman", size=14, bold=True)
    sheet["A3"].alignment = Alignment(horizontal="center", vertical="center")
    sheet["Y3"].alignment = Alignment(horizontal="center", wrap_text=True)
    sheet.row_dimensions[3].height = 30
    sheet.row_dimensions[4].height = 36
    body_font = Font(name="Times New Roman", size=10)
    wrapped = Alignment(vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=5):
        lines = 1
        for cell in row:
            cell.font = body_font
            if cell.column in {10, 12, 13, 21, 25} or (cell.column > 26 and cell.column not in money_columns):
                cell.alignment = wrapped
                width = max(widths[cell.column - 1] - 2, 1)
                lines = max(lines, sum(max(ceil(len(part) / width), 1) for part in str(cell.value or "").split("\n")))
            if cell.column in {22, 23, 26}:
                if cell.column == 23:
                    cell.number_format = "#,##0.00"
                else:
                    integer = isinstance(cell.value, (int, float)) and float(cell.value).is_integer()
                    cell.number_format = "#,##0" if integer else "#,##0.000000"
            if cell.column in money_columns:
                cell.number_format = "#,##0"
            if cell.column == 16:
                cell.number_format = "dd/mm/yyyy"
        sheet.row_dimensions[row[0].row].height = min(400, 15 * lines)
    sheet.freeze_panes = "A5"
    sheet.auto_filter.ref = f"A4:{last_column}{sheet.max_row}"
    sheet.print_title_rows = "1:4"
    sheet.print_options.horizontalCentered = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_area = f"A1:{last_column}{sheet.max_row}"
