"""Self-service fill-in for the CanBoSung sheets.

Every run lists what the reports could not fill (one row per NPP / employee / product /
unit / customer, not per order line) in ``08_BoSungDanhMuc/CanBoSung_TongHop.xlsx``.
The business fills the yellow columns of ``08_BoSungDanhMuc/BoSung_Mapping.xlsx``
(created once by the pipeline, never overwritten) and the next run applies those values
with the highest priority. Blank cells never override anything.
"""
from __future__ import annotations

import copy
import json
import os
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from typing import Any

import pandas as pd

FOLDER = "08_BoSungDanhMuc"
USER_FILE = "BoSung_Mapping.xlsx"
TODO_FILE = "CanBoSung_TongHop.xlsx"
STATE_PATH = f"{FOLDER}/_sync_state/canbosung_todo.json"
PER_PACK = "Số ĐVT nguồn trong 1 ĐVT báo cáo"
PACK_UNITS = {"thùng", "két", "bình"}
SALES_FIELDS = ("Vùng", "SS Code", "SS Name", "DB Code", "Tên NPP")
UNIT_GAP = "Thiếu quy đổi sang KÉT/THÙNG/BÌNH"

SHEETS: dict[str, dict[str, Any]] = {
    "NPP": {"key": ["DB Code"], "info": [], "values": ["Tên NPP", "SS Code", "SS Name", "Vùng"],
            "target": "npp_overrides"},
    "NhanVien": {"key": ["Mã NV"], "info": ["Tên Nhân viên"],
                 "values": ["DB Code", "Tên NPP", "SS Code", "SS Name", "Vùng"],
                 "target": "employee_overrides"},
    "SanPham": {"key": ["Mã sản phẩm"], "info": ["Tên Sản phẩm"], "values": ["Brand", "Package"],
                "target": "product_overrides"},
    "QuyDoi": {"key": ["Mã sản phẩm", "ĐVT nguồn"], "info": ["Tên Sản phẩm"],
               "values": ["ĐVT báo cáo", PER_PACK], "target": "unit_overrides"},
    "KhachHang": {"key": ["Mã Khách hàng"], "info": ["Tên Khách hàng"], "values": ["Tỉnh", "Loại KH"],
                  "target": "customer_overrides"},
}
TAIL = ["Còn thiếu", "Nguyên nhân", "Hiện có", "Gợi ý", "Số dòng", "Nguồn"]
DMS_SHEET = "SuaTrenDMS"
DMS_COLUMNS = ["Mã Đơn hàng", "Dòng nguồn", "Trường", "Lý do", "Mã SP nguồn", "ĐVT nguồn", "Số dòng", "Nguồn"]

GUIDE = [
    "Cách bổ sung dữ liệu còn thiếu cho báo cáo CTKM và báo cáo trả thưởng",
    "1. Mở file CanBoSung_TongHop.xlsx (cùng thư mục): mỗi dòng là 1 NPP / nhân viên / sản phẩm / "
    "đơn vị / khách hàng còn thiếu, cột 'Còn thiếu' ghi cần điền gì.",
    "2. Điền vào file này (BoSung_Mapping.xlsx), đúng sheet cùng tên: chép dòng từ TongHop sang, "
    "rồi điền các cột tô vàng. Ô để trống không thay đổi gì.",
    "3. NPP: điền SS Code / SS Name / Vùng / Tên NPP theo mã NPP (= mã kho xuất trên đơn). "
    "Áp dụng cho mọi đơn của NPP đó, kể cả tháng cũ.",
    "4. NhanVien: chỉ dùng khi đơn không có kho xuất, hoặc muốn ép NPP/SS cho một nhân viên.",
    "5. SanPham: Brand / Package. QuyDoi: ĐVT báo cáo (Thùng / Két / Bình) và số ĐVT nguồn trong "
    "1 ĐVT báo cáo (vd 24 chai = 1 thùng thì điền 24).",
    "6. KhachHang: Tỉnh / Loại KH theo mã khách hàng.",
    "7. Lần đồng bộ kế tiếp (tự động) sẽ áp dụng. Tháng cũ: chạy workflow với "
    "report_scope=promotion_history để xuất lại.",
    "8. Sheet SuaTrenDMS trong TongHop là lỗi dữ liệu gốc (thiếu ngày, giá...) phải sửa trên DMS.",
    "Thứ tự ưu tiên: file này > cấu hình > kho xuất của đơn > cây phòng ban hiện tại > danh mục DMS.",
]


def text(value: Any) -> str:
    if value is None or (not isinstance(value, (list, dict)) and pd.isna(value)):
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return " ".join(str(value).split())


def columns(sheet: str) -> list[str]:
    spec = SHEETS[sheet]
    return spec["key"] + spec["info"] + spec["values"] + TAIL


# --------------------------------------------------------------------------- reading
def parse_overrides(content: bytes | None) -> tuple[dict[str, dict[str, Any]], list[str]]:
    """User workbook -> override mappings; invalid rows are skipped and reported."""
    overrides: dict[str, dict[str, Any]] = {spec["target"]: {} for spec in SHEETS.values()}
    problems: list[str] = []
    if not content:
        return overrides, problems
    book = pd.read_excel(BytesIO(content), sheet_name=None, dtype=object)
    for sheet, spec in SHEETS.items():
        frame = book.get(sheet)
        if frame is None or frame.empty:
            continue
        frame.columns = [text(c) for c in frame.columns]
        missing = [c for c in spec["key"] if c not in frame.columns]
        if missing:
            problems.append(f"{sheet}: thiếu cột {missing}")
            continue
        for index, row in enumerate(frame.to_dict("records"), start=2):
            key = [text(row.get(c)) for c in spec["key"]]
            if not key[0]:
                continue
            values = {c: text(row.get(c)) for c in spec["values"] if text(row.get(c))}
            if not values:
                continue
            if sheet == "QuyDoi":
                target = values.get("ĐVT báo cáo", "")
                try:
                    per_pack = Decimal(values.get(PER_PACK, "").replace(",", "."))
                except InvalidOperation:
                    per_pack = Decimal(0)
                if target.casefold() not in PACK_UNITS or not per_pack.is_finite() or per_pack <= 0:
                    problems.append(f"QuyDoi dòng {index}: cần ĐVT báo cáo Thùng/Két/Bình và hệ số > 0")
                    continue
                overrides[spec["target"]][f"{key[0]}|{key[1]}"] = {
                    "target_unit": target, "factor": str(Decimal(1) / per_pack),
                    "units_per_pack": str(per_pack), "source": USER_FILE}
                continue
            bucket = overrides[spec["target"]].setdefault(key[0], {})
            bucket.update(values)
    return overrides, problems


def apply_overrides(cfg: dict[str, Any], overrides: dict[str, dict[str, Any]]) -> dict[str, Any]:
    result = copy.deepcopy(cfg)
    for target in ("npp_overrides", "employee_overrides", "product_overrides", "customer_overrides"):
        merged = dict(result.get(target) or {})
        for key, values in (overrides.get(target) or {}).items():
            merged[key] = {**merged.get(key, {}), **values}
        result[target] = merged
    conversions = dict(result.get("unit_conversions") or {})
    conversions.update(overrides.get("unit_overrides") or {})
    result["unit_conversions"] = conversions
    result["bosung_override_counts"] = {k: len(v) for k, v in overrides.items()}
    return result


# --------------------------------------------------------------------------- to-do list
def _known(values: dict[str, Any], fields: list[str]) -> str:
    return "; ".join(f"{f}: {text(values.get(f))}" for f in fields if text(values.get(f)))


def npp_cause(code: str, cfg: dict[str, Any]) -> str:
    """Why Tên NPP / SS of a warehouse unit could not be derived from the department tree."""
    if code not in cfg.get("npp_units", {}):
        return "Mã NPP (kho xuất) không có trên cây phòng ban DMS hiện tại (đơn vị đã xóa/đổi mã)"
    province = code.split("-")[1] if code.count("-") == 2 else ""
    candidates = [c for c in cfg.get("supervisor_candidates", {}).get(province, "").split("; ") if c]
    if not candidates:
        return (f"NPP chưa gán nhân viên chức vụ 'Giám sát kinh doanh' và tỉnh {province} "
                "không có giám sát nào trên DMS")
    return (f"NPP chưa gán giám sát; tỉnh {province} có {len(candidates)} giám sát "
            "nên không tự chọn được")


def customer_cause(code: str, field: str, cfg: dict[str, Any]) -> str:
    matches = [m for m in cfg.get("customer_catalogue", {}).values() if m.get("customer_code") == code]
    if not matches:
        return "Khách không còn trong danh mục khách hàng DMS (đã xóa hoặc đổi mã/ID)"
    source = {"Tỉnh": "tinh_thanh_moi", "Loại KH": "loai_kh"}.get(field, field)
    return f"Danh mục khách hàng DMS để trống {field} ({source})"


def product_cause(sku: str, field: str, cfg: dict[str, Any]) -> str:
    if field == "Package":
        return ("DMS không có trường Package; danh mục Package (Danh Muc San Pham.xlsx) chưa có SP này"
                if sku not in cfg.get("products", {}) or "Brand" in cfg["products"][sku]
                else "Danh mục sản phẩm cấu hình thiếu Package")
    return "Danh mục sản phẩm DMS để trống nhãn hiệu (nhan_hieu)"


def unit_cause(sku: str, unit: str) -> str:
    if not unit:
        return "Dòng đơn hàng trên DMS không ghi ĐVT"
    return f"SP chưa khai quy cách {unit} → Thùng/Két/Bình (dvt_chan/dvt_le/hsqd) trên danh mục DMS"


def todo_rows(issues: pd.DataFrame | None, cfg: dict[str, Any]) -> list[dict[str, Any]]:
    """One CanBoSung issue -> one fill-in row (not yet aggregated)."""
    rows: list[dict[str, Any]] = []
    if issues is None or issues.empty:
        return rows
    for rec in issues.to_dict("records"):
        field, reason = text(rec.get("Trường")), text(rec.get("Lý do"))
        sku, unit = text(rec.get("Mã SP nguồn")), text(rec.get("ĐVT nguồn"))
        if reason.startswith(UNIT_GAP):
            rows.append({"_sheet": "QuyDoi", "Mã sản phẩm": sku, "ĐVT nguồn": unit,
                         "Tên Sản phẩm": text(rec.get("Tên SP nguồn")), "Còn thiếu": ["Quy đổi"],
                         "Nguyên nhân": unit_cause(sku, unit),
                         "Gợi ý": "" if unit else "Đơn hàng thiếu ĐVT: nên sửa đơn trên DMS; "
                                                  "điền ở đây sẽ áp dụng cho dòng không ĐVT của SP này"})
        elif field in SALES_FIELDS:
            code = text(rec.get("Kho/NPP"))
            if code:
                known = {"DB Code": code, **cfg.get("npp_units", {}).get(code, {}),
                         **cfg.get("npp_overrides", {}).get(code, {})}
                hint = cfg.get("supervisor_candidates", {}).get(code.split("-")[1], "") if "-" in code else ""
                rows.append({"_sheet": "NPP", "DB Code": code, "Còn thiếu": [field],
                             "Nguyên nhân": npp_cause(code, cfg),
                             "Hiện có": _known(known, ["Tên NPP", "SS Code", "SS Name", "Vùng"]),
                             "Gợi ý": f"Giám sát cùng tỉnh: {hint}" if hint and field.startswith("SS") else ""})
            else:
                employee = text(rec.get("Mã NV"))
                rows.append({"_sheet": "NhanVien", "Mã NV": employee,
                             "Tên Nhân viên": text(rec.get("Tên Nhân viên")), "Còn thiếu": [field],
                             "Nguyên nhân": "Đơn không có kho xuất và nhân viên không thuộc NPP nào "
                                            "trên cây phòng ban hiện tại",
                             "Hiện có": _known(cfg.get("employees", {}).get(employee, {}), list(SALES_FIELDS)),
                             "Gợi ý": "Đơn không có kho xuất"})
        elif field in ("Brand", "Package"):
            rows.append({"_sheet": "SanPham", "Mã sản phẩm": sku, "Tên Sản phẩm": text(rec.get("Tên SP nguồn")),
                         "Còn thiếu": [field], "Nguyên nhân": product_cause(sku, field, cfg),
                         "Hiện có": _known(cfg.get("products", {}).get(sku, {}), ["Brand", "Package"])})
        elif field in ("Tỉnh", "Loại KH"):
            rows.append({"_sheet": "KhachHang", "Mã Khách hàng": text(rec.get("Mã Khách hàng")),
                         "Tên Khách hàng": text(rec.get("Tên Khách hàng")), "Còn thiếu": [field],
                         "Nguyên nhân": customer_cause(text(rec.get("Mã Khách hàng")), field, cfg)})
        else:
            rows.append({"_sheet": DMS_SHEET, "Mã Đơn hàng": text(rec.get("Mã Đơn hàng")),
                         "Dòng nguồn": text(rec.get("Dòng nguồn")), "Trường": field, "Lý do": reason,
                         "Mã SP nguồn": sku, "ĐVT nguồn": unit})
    return rows


def _row_key(row: dict[str, Any]) -> tuple[str, ...]:
    sheet = row["_sheet"]
    names = SHEETS[sheet]["key"] if sheet in SHEETS else ["Mã Đơn hàng", "Dòng nguồn", "Trường"]
    return (sheet, *(text(row.get(n)) for n in names))


def aggregate(rows: list[dict[str, Any]], label: str | None = None) -> list[dict[str, Any]]:
    """Merge rows with the same sheet+key: union of missing fields and sources, summed counts."""
    merged: dict[tuple[str, ...], dict[str, Any]] = {}
    for row in rows:
        key = _row_key(row)
        count = int(row.get("Số dòng") or 1)
        sources = set(row.get("Nguồn") or ([label] if label else []))
        if key not in merged:
            merged[key] = {**row, "Còn thiếu": sorted(set(row.get("Còn thiếu") or [])),
                           "Số dòng": count, "Nguồn": sorted(sources)}
            continue
        target = merged[key]
        target["Còn thiếu"] = sorted(set(target["Còn thiếu"]) | set(row.get("Còn thiếu") or []))
        target["Số dòng"] += count
        target["Nguồn"] = sorted(set(target["Nguồn"]) | sources)
        for column, value in row.items():
            if column not in {"Còn thiếu", "Số dòng", "Nguồn"} and not text(target.get(column)) and text(value):
                target[column] = value
    return list(merged.values())


def _filled(row: dict[str, Any], overrides: dict[str, dict[str, Any]]) -> bool:
    sheet = row["_sheet"]
    if sheet not in SHEETS:
        return False
    target = overrides.get(SHEETS[sheet]["target"], {})
    if sheet == "QuyDoi":
        return f"{text(row.get('Mã sản phẩm'))}|{text(row.get('ĐVT nguồn'))}" in target
    values = target.get(text(row.get(SHEETS[sheet]["key"][0])), {})
    return all(text(values.get(field)) for field in row.get("Còn thiếu") or [])


def todo_frames(state: dict[str, list[dict[str, Any]]],
                overrides: dict[str, dict[str, Any]] | None = None) -> dict[str, pd.DataFrame]:
    rows = aggregate([row for label_rows in state.values() for row in label_rows])
    rows = [row for row in rows if not _filled(row, overrides or {})]
    frames: dict[str, pd.DataFrame] = {}
    for sheet in [*SHEETS, DMS_SHEET]:
        wanted = columns(sheet) if sheet in SHEETS else DMS_COLUMNS
        records = [{c: (", ".join(row[c]) if isinstance(row.get(c), list) else row.get(c, ""))
                    for c in wanted} for row in rows if row["_sheet"] == sheet]
        frame = pd.DataFrame(records, columns=wanted, dtype=object)
        if not frame.empty:
            frame = frame.sort_values("Số dòng", ascending=False, kind="stable")
        frames[sheet] = frame.reset_index(drop=True)
    return frames


def summary(frames: dict[str, pd.DataFrame]) -> dict[str, Any]:
    """Counts only (no names): keys per sheet and which fields they miss."""
    result: dict[str, Any] = {}
    for sheet, frame in frames.items():
        if frame.empty:
            continue
        if sheet in SHEETS:
            missing: dict[str, int] = {}
            for value in frame["Còn thiếu"]:
                for field in str(value).split(", "):
                    missing[field] = missing.get(field, 0) + 1
            causes = frame["Nguyên nhân"].astype(str).value_counts()
            result[sheet] = {"keys": len(frame), "lines": int(frame["Số dòng"].sum()), "missing": missing,
                             "causes": {k[:90]: int(v) for k, v in causes.items()}}
            if sheet in {"NPP", "SanPham", "QuyDoi"}:
                key = SHEETS[sheet]["key"][0]
                result[sheet]["items"] = [f"{r[key]}|{r['Số dòng']}|{str(r['Nguyên nhân'])[:60]}"
                                          for r in frame.head(10).to_dict("records")]
        else:
            reasons = (frame["Trường"].astype(str) + ": " + frame["Lý do"].astype(str)).value_counts()
            result[sheet] = {"rows": len(frame), "reasons": {k: int(v) for k, v in reasons.head(5).items()}}
    return result


# --------------------------------------------------------------------------- writing
def write_book(frames: dict[str, pd.DataFrame], path: Path, guide: list[str] | None = None) -> Path:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    path.parent.mkdir(parents=True, exist_ok=True)
    yellow, grey = PatternFill("solid", fgColor="FFF2CC"), PatternFill("solid", fgColor="D9D9D9")
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        if guide:
            pd.DataFrame({"Hướng dẫn": guide}).to_excel(writer, sheet_name="HuongDan", index=False)
            writer.sheets["HuongDan"].column_dimensions["A"].width = 140
        for sheet, frame in frames.items():
            frame.to_excel(writer, sheet_name=sheet, index=False)
            ws = writer.sheets[sheet]
            values = set(SHEETS[sheet]["values"]) if sheet in SHEETS else set()
            for index, column in enumerate(frame.columns, start=1):
                cell = ws.cell(row=1, column=index)
                cell.font = Font(bold=True)
                cell.alignment = Alignment(wrap_text=True, vertical="center")
                cell.fill = yellow if column in values else grey
                if column in values:
                    for row in range(2, len(frame) + 2):
                        ws.cell(row=row, column=index).fill = yellow
                width = 40 if column in {"Hiện có", "Gợi ý", "Lý do"} else 18
                ws.column_dimensions[get_column_letter(index)].width = width
            ws.freeze_panes = "B2"
    return path


def user_template(frames: dict[str, pd.DataFrame]) -> dict[str, pd.DataFrame]:
    return {sheet: frames.get(sheet, pd.DataFrame(columns=columns(sheet))) for sheet in SHEETS}


# --------------------------------------------------------------------------- SharePoint
def _notice(level: str, title: str, message: str) -> None:
    if os.environ.get("GITHUB_ACTIONS") == "true":
        print(f"::{level} title={title}::{' '.join(message.split())[:3000]}")


def sharepoint_context(sharepoint: Any = None, drive: str = "") -> tuple[Any, str]:
    if sharepoint is None:
        from sharepoint_semantic import SemanticSharePointClient

        sharepoint = SemanticSharePointClient.from_env()
    if not drive:
        drive = os.environ.get("SHAREPOINT_DRIVE_ID", "").strip() or sharepoint.get_drive_id(
            sharepoint.get_site_id())
    return sharepoint, drive


def load_overrides(sharepoint: Any = None, drive: str = "") -> dict[str, dict[str, Any]]:
    """Read the user's workbook; any failure leaves the reports exactly as without it."""
    try:
        sharepoint, drive = sharepoint_context(sharepoint, drive)
        content = sharepoint.download_file_bytes(drive, f"{FOLDER}/{USER_FILE}")
        overrides, problems = parse_overrides(content)
    except Exception as exc:
        _notice("warning", "BoSung_Mapping not applied", f"{type(exc).__name__}: {exc}")
        return {}
    counts = {k: len(v) for k, v in overrides.items() if v}
    _notice("notice", "BoSung_Mapping", f"applied={counts} invalid_rows={len(problems)} "
                                        f"{'; '.join(problems[:5])}")
    return overrides


def publish(updates: dict[str, list[dict[str, Any]]], sharepoint: Any = None, drive: str = "",
            dry_run: bool = False, output_dir: Path = Path("output")) -> dict[str, Any]:
    """Merge this run's per-month rows into the shared state and refresh the to-do workbook.

    ``updates`` maps a source label (e.g. "CTKM 2026-10") to its aggregated rows; a label
    with no rows clears that month. Never raises: the reports are already built.
    """
    result: dict[str, Any] = {"labels": sorted(updates)}
    try:
        if dry_run and sharepoint is None:
            state, overrides = {}, {}
        else:
            sharepoint, drive = sharepoint_context(sharepoint, drive)
            state = sharepoint.download_json(drive, STATE_PATH) or {}
            overrides, _ = parse_overrides(sharepoint.download_file_bytes(drive, f"{FOLDER}/{USER_FILE}"))
        state = {k: v for k, v in state.items() if isinstance(v, list)}
        state.update(updates)
        frames = todo_frames(state, overrides)
        result["todo"] = summary(frames)
        todo = write_book(frames, output_dir / TODO_FILE, GUIDE)
        _notice("notice", "CanBoSung tổng hợp", json.dumps(result["todo"], ensure_ascii=False)
                + f" sources={sorted(state)}")
        if dry_run:
            return result
        sharepoint.upload_json(drive, STATE_PATH, state)
        if not sharepoint.get_item_by_path(drive, f"{FOLDER}/{USER_FILE}"):
            template = write_book(user_template(frames), output_dir / USER_FILE, GUIDE)
            sharepoint.upload_file(drive, template, FOLDER)
            result["user_file_created"] = True
        try:
            sharepoint.upload_file(drive, todo, FOLDER)
            result["todo_published"] = True
        except Exception as exc:  # open in Excel (423): refreshed next run
            _notice("warning", "CanBoSung_TongHop not updated", f"{type(exc).__name__}: {exc}")
    except Exception as exc:
        result["error"] = f"{type(exc).__name__}: {exc}"
        _notice("warning", "CanBoSung tổng hợp failed", result["error"])
    return result
