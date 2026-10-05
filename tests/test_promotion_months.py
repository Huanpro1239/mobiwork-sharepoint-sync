from __future__ import annotations

import tempfile
import os
import json
from io import BytesIO
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import Mock, patch

import pandas as pd
from openpyxl import load_workbook

from mobiwork import ReportConfig
from promotion_detail import build_report, unit_trace
from promotion_months import discover_bill_months, select_order_month
from promotion_workbook import write_detail_workbook
from test_promotion_detail import config, source
import promotion_detail as module


class PromotionMonthTests(unittest.TestCase):
    def test_history_dry_reads_all_masters_and_production_prepares_before_writes(self):
        for dry in (True, False):
            with self.subTest(dry=dry), tempfile.TemporaryDirectory() as folder:
                cfg = config()
                cfg.update(publish_enabled=True, allow_incomplete_publish=True)
                buffers = []
                for month in (7, 9):
                    buffer = BytesIO()
                    row = source(ngay_dat=f"2026-{month:02}-09", gia_truoc_vat=None if month == 9 and not dry else 100)
                    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
                        pd.DataFrame([row]).to_excel(writer, sheet_name="ChiTietSP", index=False)
                    buffers.append(buffer.getvalue())
                sp = Mock()
                sp.download_file_bytes.side_effect = buffers
                bill = ReportConfig(key="bill", enabled=True, name="DonBanHang", folder="04_DonBanHang")
                cwd = Path.cwd()
                try:
                    os.chdir(folder)
                    with patch.object(module, "load_config", return_value=cfg), \
                         patch.object(module, "load_reports", return_value=[bill]), \
                         patch.object(module, "discover_bill_months", return_value=[date(2026, 7, 1), date(2026, 9, 1)]), \
                         patch.object(module, "write_detail_workbook", side_effect=[Path("july.xlsx"), Path("september.xlsx")]), \
                         patch.object(module.SemanticSharePointClient, "from_env", return_value=sp), \
                         patch.dict(os.environ, {"PROMOTION_DETAIL_SCOPE": "all_existing", "DRY_RUN": "true" if dry else "false", "SHAREPOINT_DRIVE_ID": "drive"}):
                        if dry:
                            manifest = module.run()
                            self.assertEqual(manifest["source_months"], ["2026-07", "2026-09"])
                            self.assertTrue(all(r["source_scope"] == "monthly_master" for r in manifest["results"]))
                            self.assertEqual([r["rows"] for r in manifest["results"]], [1, 1])
                        else:
                            with self.assertRaisesRegex(ValueError, "Nothing published"):
                                module.run()
                            manifest = json.loads(Path("output/promotion_detail_manifest.json").read_text())
                            self.assertEqual(manifest["results"][1]["blocking_issues"], 1)
                        sp.upload_file.assert_not_called()
                        self.assertEqual(sp.download_file_bytes.call_count, 2)
                finally:
                    os.chdir(cwd)

    def test_discovery_lists_existing_masters_across_years_without_future(self):
        sp = Mock()
        sp.list_folder_children.side_effect = [
            [{"name": "2025", "folder": {}}, {"name": "2026", "folder": {}},
             {"name": "_sync_state", "folder": {}}],
            [{"name": "12", "folder": {}}],
            [{"name": "07", "folder": {}}, {"name": "09", "folder": {}},
             {"name": "11", "folder": {}}, {"name": "readme.txt", "file": {}}]]
        sp.get_item_by_path.return_value = {"file": {}}
        bill = ReportConfig(key="bill", enabled=True, name="DonBanHang", folder="04_DonBanHang")
        self.assertEqual(discover_bill_months(sp, "drive", bill, date(2026, 10, 5)),
                         [date(2025, 12, 1), date(2026, 7, 1), date(2026, 9, 1)])
        self.assertEqual(sp.get_item_by_path.call_count, 3)

    def test_missing_master_or_empty_history_fails(self):
        bill = ReportConfig(key="bill", enabled=True, name="DonBanHang", folder="04_DonBanHang")
        sp = Mock()
        sp.list_folder_children.side_effect = [[{"name": "2026", "folder": {}}],
                                              [{"name": "07", "folder": {}}]]
        sp.get_item_by_path.return_value = None
        with self.assertRaisesRegex(ValueError, "master is missing"):
            discover_bill_months(sp, "drive", bill, date(2026, 10, 5))
        sp.list_folder_children.side_effect = [[]]
        with self.assertRaisesRegex(ValueError, "No historical"):
            discover_bill_months(sp, "drive", bill, date(2026, 10, 5))

    def test_order_month_filter_keeps_bad_dates_for_validation(self):
        detail = pd.DataFrame([source(ngay_dat="2026-07-09"),
                               source(stt=2, ngay_dat="2026-09-03"),
                               source(stt=3, ngay_dat="bad"),
                               source(stt=4, ngay_dat=None, ngay_ban_hang=None)])
        selected, excluded = select_order_month(detail, date(2026, 7, 1))
        self.assertEqual(excluded, 1)
        self.assertEqual(selected["stt"].tolist(), [1, 3, 4])
        self.assertTrue(select_order_month(pd.DataFrame(), date(2026, 7, 1))[0].empty)

    def test_workbook_template_headers_typed_values_and_literal_data(self):
        detail = pd.DataFrame([source(ten_kh="=literal")])
        report, issues = build_report(detail, config())
        frames = {"BaoCao": report, "CanBoSung": issues, "DonViTinh": unit_trace(detail, report, config())}
        with tempfile.TemporaryDirectory() as folder:
            path = write_detail_workbook(frames, "report.xlsx", date(2026, 10, 1), Path(folder))
            workbook = load_workbook(path)
            sheet = workbook["BaoCao"]
            self.assertEqual(sheet["B2"].value, "2026-10")
            self.assertEqual(sheet["A4"].value, "Vùng")
            self.assertEqual(sheet["L5"].value, "=literal")
            self.assertEqual(sheet["L5"].data_type, "s")
            self.assertEqual(sheet["W5"].value, 1000)
            self.assertEqual(sheet["Z5"].value, 0)
            self.assertEqual(sheet["Z5"].number_format, "#,##0")
            self.assertEqual(sheet["W5"].number_format, "#,##0.00")
            self.assertEqual(sheet["P5"].data_type, "d")
            self.assertEqual(sheet.freeze_panes, "A5")
            self.assertEqual(sheet.auto_filter.ref, "A4:Z5")
            self.assertEqual(set(workbook.sheetnames), set(frames))
            workbook.close()
