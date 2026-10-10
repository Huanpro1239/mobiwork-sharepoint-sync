from __future__ import annotations

import copy
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path
from unittest.mock import Mock

import pandas as pd
from openpyxl import load_workbook

import promotion_bonus_calc as calc
from promotion_detail import COLUMNS
from promotion_tracking import BASE_COLUMNS, tracking_report
from promotion_workbook import write_detail_workbook
from test_promotion_bonus_calc import C1, bill, qty_program

JUL, AUG, SEP = date(2026, 7, 1), date(2026, 8, 1), date(2026, 9, 1)


def epoch(day):
    return int(datetime(day.year, day.month, day.day, tzinfo=calc.VN_TZ).timestamp() * 1000)


def source(month, quantity, **extra):
    return bill([{"so_luong": quantity, "thanh_tien": quantity * 100,
                  "ngay_giao_hang": f"{month:%Y-%m}-10T03:00:00Z", **extra}])


def build(mode="monthly", amounts=(40, 40, 80), multiple=False, missing=None, extra=None):
    program = qty_program(minimum=72, customers=(C1,), multiple=multiple)
    program.update(startDate=epoch(JUL), endDate=epoch(date(2026, 9, 30)))
    months = {m: source(m, q, **(extra or {})) for m, q in zip((JUL, AUG, SEP), amounts, strict=True)}
    if missing:
        months[missing] = None
    current_lines = calc.sold_lines(months[SEP], SEP, date(2026, 9, 30))
    lines = current_lines if mode == "monthly" else [line for m, frame in months.items() if frame is not None
                                                     for line in calc.sold_lines(frame, m, date(m.year, m.month, 30))]
    results, _ = calc.compute([program], lines, {})
    loader = Mock(side_effect=lambda month: months.get(month))
    frame = tracking_report(results, SEP, date(2026, 9, 30), loader, {}, {}, {program["_id"]: mode})
    return frame, loader


class TrackingTests(unittest.TestCase):
    def test_monthly_does_not_use_accumulated_sales_to_grant_slots(self):
        frame, loader = build()
        row = frame.iloc[0]
        self.assertEqual(list(frame.columns[:20]), BASE_COLUMNS)
        self.assertEqual([row[f"TỔNG TÍCH LŨY T{m:02d}/2026"] for m in (7, 8, 9)], [40, 40, 80])
        self.assertEqual([row[f"SỐ SUẤT ĐẠT T{m:02d}/2026"] for m in (7, 8, 9)], [0, 0, 1])
        self.assertEqual((row["TỔNG TÍCH LŨY"], row["Thực hiện kỳ xét thưởng"], row["CÒN LẠI"]), (160, 80, 0))
        self.assertEqual(row["SỐ SUẤT"], 1)
        self.assertIsNone(row["NGÀY ĐĂNG KÝ"])
        # Jul–Sep plus the June master (orders created in June, delivered in July), once each.
        self.assertEqual(sorted(call.args[0] for call in loader.call_args_list), [date(2026, 6, 1), JUL, AUG, SEP])

    def test_cumulative_slots_are_increments_not_repeated_reward(self):
        row = build("cumulative", (40, 40, 80), multiple=True)[0].iloc[0]
        self.assertEqual([row[f"SỐ SUẤT ĐẠT T{m:02d}/2026"] for m in (7, 8, 9)], [0, 1, 1])
        self.assertEqual(row["TỔNG SỐ SUẤT ĐẠT"], 2)

    def test_missing_month_is_blank_and_total_unavailable_not_zero(self):
        frame, _ = build(missing=AUG)
        row = frame.iloc[0]
        self.assertIsNone(row["TỔNG TÍCH LŨY T08/2026"])
        self.assertIsNone(row["TỔNG TÍCH LŨY"])
        self.assertIsNone(row["TỔNG SỐ SUẤT ĐẠT"])
        self.assertEqual(row["SỐ SUẤT ĐẠT T09/2026"], 1)
        self.assertIn("Thiếu đơn hàng 08/2026", row["Thông tin nguồn"])

    def test_monthly_remaining_uses_current_month_even_when_prior_month_reached(self):
        row = build(amounts=(80, 80, 10))[0].iloc[0]
        self.assertEqual(row["CÒN LẠI"], -62)
        self.assertEqual(row["TỔNG SỐ SUẤT ĐẠT"], 2)
        self.assertEqual(row["Trạng thái trả thưởng"], "Không")

    def test_thresholds_use_source_precision_before_display_rounding(self):
        row = build(amounts=(71.999, 0.0009, 71.999))[0].iloc[0]
        self.assertEqual(row["SỐ SUẤT ĐẠT T07/2026"], 0)
        self.assertEqual(row["SỐ SUẤT ĐẠT T09/2026"], 0)
        self.assertAlmostEqual(row["CÒN LẠI"], -0.001)
        row = build("cumulative", amounts=(71.999, 0.0009, 0))[0].iloc[0]
        self.assertEqual(row["TỔNG SỐ SUẤT ĐẠT"], 0)
        self.assertAlmostEqual(row["TỔNG TÍCH LŨY"], 71.9999)

    def test_gifts_and_wrong_units_never_inflate_progress(self):
        row = build(extra={"ten_dvt": "Thùng", "ma_dvt": "Thùng"})[0].iloc[0]
        self.assertEqual(row["TỔNG TÍCH LŨY"], 0)
        row = build(extra={"is_km": True})[0].iloc[0]
        self.assertEqual(row["TỔNG TÍCH LŨY"], 0)

    def test_future_months_stay_blank_and_unfetched(self):
        program = qty_program(customers=(C1,))
        program.update(startDate=epoch(JUL), endDate=epoch(date(2026, 9, 30)))
        detail = source(JUL, 80)
        results, _ = calc.compute([program], calc.sold_lines(detail, JUL, date(2026, 7, 31)), {})
        original = copy.deepcopy(results)
        loader = Mock(return_value=detail)
        frame = tracking_report(results, JUL, date(2026, 7, 31), loader, {}, {}, {program["_id"]: "monthly"})
        self.assertIsNone(frame.iloc[0]["TỔNG TÍCH LŨY T08/2026"])
        self.assertEqual(frame.iloc[0]["TỔNG TÍCH LŨY"], 80)
        self.assertEqual(sorted(call.args[0] for call in loader.call_args_list), [date(2026, 6, 1), JUL])
        self.assertEqual(results, original)

    def test_historical_missing_unit_marks_affected_period_unavailable(self):
        program = qty_program(customers=(C1,))
        program.update(startDate=epoch(JUL), endDate=epoch(date(2026, 9, 30)))
        source_sep = source(SEP, 80)
        results, _ = calc.compute([program], calc.sold_lines(source_sep, SEP, date(2026, 9, 30)), {})
        frame = tracking_report(results, SEP, date(2026, 9, 30),
                                lambda month: source_sep if month == SEP else source(month, 40, ten_dvt="", ma_dvt=""),
                                {}, {}, {program["_id"]: "monthly"})
        self.assertIsNone(frame.iloc[0]["TỔNG TÍCH LŨY"])
        self.assertIn("Thiếu ĐVT", frame.iloc[0]["Thông tin nguồn"])

    def test_cumulative_upper_bound_and_returns_reconcile_monthly_slots(self):
        program = qty_program(minimum=72, maximum=100, customers=(C1,))
        program.update(startDate=epoch(JUL), endDate=epoch(date(2026, 9, 30)))
        months = {m: source(m, q) for m, q in zip((JUL, AUG, SEP), (80, 30, -25), strict=True)}
        lines = [line for m, frame in months.items() for line in calc.sold_lines(frame, m, date(m.year, m.month, 30))]
        results, _ = calc.compute([program], lines, {})
        frame = tracking_report(results, SEP, date(2026, 9, 30), months.get, {}, {}, {program["_id"]: "cumulative"})
        row = frame.iloc[0]
        self.assertEqual([row[f"SỐ SUẤT ĐẠT T{m:02d}/2026"] for m in (7, 8, 9)], [1, -1, 1])
        self.assertEqual(row["TỔNG SỐ SUẤT ĐẠT"], calc.reward_multiplier(85, calc.parse_rule(program)))

    def test_workbook_places_two_template_views_first_with_original_detail_columns(self):
        tracking = build()[0]
        detail = pd.DataFrame([{**dict.fromkeys(COLUMNS), "THÀNH TIỀN": 1200.25,
                                "Số lượng Khuyến mãi": 2, "Tên Khách hàng": "=literal"}])
        frames = {"TheoDoiTichLuy": tracking, "KhuyenMaiDonHang": detail, "BaoCao": detail}
        with tempfile.TemporaryDirectory() as folder:
            path = write_detail_workbook(frames, "report.xlsx", SEP, Path(folder))
            book = load_workbook(path)
            try:
                self.assertEqual(book.sheetnames[:2], ["TheoDoiTichLuy", "KhuyenMaiDonHang"])
                self.assertEqual(book.active.title, "TheoDoiTichLuy")
                self.assertEqual(book["TheoDoiTichLuy"]["G5"].value, "Mã KH")
                self.assertEqual(book["TheoDoiTichLuy"].freeze_panes, "H6")
                self.assertEqual(book["TheoDoiTichLuy"]["Q6"].value, 1)
                self.assertEqual(book["TheoDoiTichLuy"]["P6"].data_type, "n")
                invoice = book["KhuyenMaiDonHang"]
                self.assertEqual(invoice["W5"].value, 1200.25)
                self.assertEqual(invoice["Z5"].value, 2)
                self.assertEqual(invoice["L5"].data_type, "s")
                self.assertEqual(invoice.freeze_panes, "A5")
            finally:
                book.close()

    def test_delivery_month_counts_orders_created_in_the_previous_month(self):
        program = qty_program(minimum=72, customers=(C1,))
        program.update(startDate=epoch(AUG), endDate=epoch(date(2026, 9, 30)))
        # Created in August (August master), delivered on 2 September.
        months = {AUG: bill([{"so_luong": 50, "ma_phieu": "DH8", "ngay_giao_hang": "2026-09-02T03:00:00Z"}]),
                  SEP: bill([{"so_luong": 30, "ma_phieu": "DH9", "ngay_giao_hang": "2026-09-10T03:00:00Z"}])}
        lines, missing, neighbours = calc.delivered_lines(months.get, SEP, date(2026, 9, 30), newest=SEP)
        self.assertEqual(sorted(line["raw"]["ma_phieu"] for line in lines), ["DH8", "DH9"])
        self.assertEqual((missing, neighbours), (set(), set()))
        results, _ = calc.compute([program], lines, {})
        frame = tracking_report(results, SEP, date(2026, 9, 30), months.get, {}, {}, {program["_id"]: "monthly"})
        row = frame.iloc[0]
        self.assertEqual((row["TỔNG TÍCH LŨY T08/2026"], row["TỔNG TÍCH LŨY T09/2026"]), (0, 80))
        self.assertEqual(row["SỐ SUẤT ĐẠT T09/2026"], 1)


if __name__ == "__main__":
    unittest.main()
