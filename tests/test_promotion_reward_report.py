from __future__ import annotations

import copy
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

import pandas as pd

import promotion_bonus as bonus
import promotion_bonus_calc as calc
from promotion_bonus_ui import target_cells
from promotion_reward_report import invoice_promotions, program_coverage, reward_ledger
from test_promotion_bonus_calc import C1, OCT1, OCT31, amount_program, bill, qty_program


def cash_program():
    program = amount_program(minimum=300_000, maximum=0)
    program["ptype"] = {"value": calc.MONEY_TYPE}
    program["products"][0]["khuyen_mai"] = "550000"
    program["products"][0]["san_pham_mua"][0]["ma_san_pham"] = "230100110"
    return program


class RewardCoverageTests(unittest.TestCase):
    def test_same_pool_tiers_choose_one_reward_at_each_boundary(self):
        program = qty_program(minimum=16, maximum=30, units=("Thùng",), customers=(C1,), multiple=True)
        program["products"][0]["san_pham_khuyen_mai"][0][0]["so_luong"] = 1
        second = copy.deepcopy(program["products"][0])
        second["yeu_cau"] = {"qualityMin": 30, "qualityMax": 0}
        second["san_pham_khuyen_mai"][0][0]["so_luong"] = 2
        program["products"].append(second)
        rule = calc.parse_rule(program)
        self.assertEqual(len(rule.tiers), 2)
        for qty, expected in [(15, 0), (16, 1), (29, 1), (30, 2), (59, 2), (60, 4)]:
            with self.subTest(qty=qty):
                source = bill([{"so_luong": qty, "ten_dvt": "Thùng", "ma_dvt": "Thùng"}])
                results, issues = calc.compute([program], calc.sold_lines(source, OCT1, OCT31), {})
                self.assertEqual(issues, [])
                self.assertEqual(len(results[0].rows), 1)
                ledger = reward_ledger(results, OCT1)
                self.assertEqual(len(ledger), 1)
                self.assertEqual(ledger["Số lượng quà đủ điều kiện"].sum(), expected)
                self.assertEqual(results[0].rows[0]["objThucHien"][calc.TARGET_ID], qty)
                if expected:
                    self.assertEqual(target_cells(results[0].rows[0], results[0].targets)[0]["Tỷ lệ (%)"], 100)

    def test_tiers_with_different_pools_or_overlapping_ranges_are_not_guessed(self):
        for different_pool in (True, False):
            program = qty_program(minimum=16, maximum=30)
            second = copy.deepcopy(program["products"][0])
            if different_pool:
                second["san_pham_mua"][0]["ma_san_pham"] = "OTHER"
                second["yeu_cau"] = {"qualityMin": 30, "qualityMax": 0}
            program["products"].append(second)
            with self.subTest(different_pool=different_pool), self.assertRaises(ValueError):
                calc.parse_rule(program)

    def test_disconnected_tier_ranges_are_kept_unsupported(self):
        program = qty_program(minimum=16, maximum=30)
        second = copy.deepcopy(program["products"][0])
        second["yeu_cau"] = {"qualityMin": 40, "qualityMax": 0}
        program["products"].append(second)
        with self.assertRaisesRegex(ValueError, "gaps"):
            calc.parse_rule(program)

    def test_cash_is_never_a_physical_quantity_and_allocation_reconciles(self):
        source = bill([{"so_luong": 72, "thanh_tien": 400_000}])
        results, _ = calc.compute([cash_program()], calc.sold_lines(source, OCT1, OCT31), {})
        report, _ = bonus._template_report(Mock(), {}, results, OCT1, detail_config={})
        gift = report[report["Mã Đơn hàng"] == calc.GIFT_ORDER].iloc[0]
        self.assertEqual(gift["Số lượng Khuyến mãi"], 0)
        self.assertEqual(gift["Tiền mặt trả thưởng (đ)"], 550_000)
        self.assertEqual(report["Tiền mặt phân bổ (đ)"].sum(), 550_000)
        ledger = reward_ledger(results, OCT1)
        self.assertEqual(ledger["Số lượng quà đủ điều kiện"].sum(), 0)
        self.assertEqual(ledger["Tiền mặt đủ điều kiện (đ)"].sum(), 550_000)

    def test_missing_gift_price_does_not_hide_known_cash(self):
        source = bill([{"so_luong": 72, "thanh_tien": 400_000}])
        results, _ = calc.compute([qty_program(customers=(C1,))], calc.sold_lines(source, OCT1, OCT31), {})
        results[0].rows[0]["_rewards"].append((calc.MONEY_SKU, calc.MONEY_NAME, calc.MONEY_UNIT, 550_000))
        frame = calc.detail_source(results, "10/2026")[0][1]
        self.assertEqual(frame.iloc[0]["_cash_alloc"], 550_000)
        self.assertTrue(pd.isna(frame.iloc[0]["_gift_alloc"]))
        self.assertTrue(pd.isna(frame.iloc[0][calc.BONUS_VALUE]))
        report, _ = bonus._template_report(Mock(), {}, results, OCT1, detail_config={})
        orders = bonus.reward_by_order(report)
        self.assertEqual(orders["Tiền mặt phân bổ (đ)"].sum(), 550_000)
        self.assertTrue(orders["Giá trị quà ước tính (đ)"].isna().all())

    def test_held_cash_remains_visible_without_becoming_payable(self):
        source = bill([{"so_luong": 72, "thanh_tien": 400_000}])
        results, _ = calc.compute([cash_program()], calc.sold_lines(source, OCT1, OCT31), {})
        calc.hold_rewards(results, "Thiếu dữ liệu kỳ tích lũy: 2026-04")
        ledger = reward_ledger(results, OCT1)
        self.assertEqual(ledger["Tiền mặt dự kiến (đ)"].sum(), 550_000)
        self.assertEqual(ledger["Tiền mặt đủ điều kiện (đ)"].sum(), 0)
        self.assertIn("2026-04", ledger.iloc[0]["Trạng thái"])
        coverage = program_coverage(results, OCT1, [])
        self.assertEqual(coverage.iloc[0]["Khách cần xem"], 1)
        self.assertEqual(coverage.iloc[0]["Loại thưởng"], "Tiền mặt")

    def test_every_program_level_survives_empty_sales_and_unsupported_rules(self):
        zero = qty_program(pid="empty", customers=())
        unsupported = qty_program(pid="unsupported", customers=(C1,))
        unsupported["ptype"] = {"value": "UNKNOWN"}
        results, issues = calc.compute([zero, unsupported], [], {})
        coverage = program_coverage(results, OCT1, issues)
        self.assertEqual(set(coverage["ID CT"]), {"empty", "unsupported"})
        self.assertEqual(coverage.iloc[0]["Trạng thái"], "Không có khách đăng ký")
        self.assertEqual(coverage.iloc[1]["Khách đăng ký"], 1)
        self.assertEqual(coverage.iloc[1]["Trạng thái"], "Chưa tính được")
        self.assertIn("UNKNOWN", coverage.iloc[1]["Lý do"])

    def test_actual_promotions_keep_one_sale_and_each_gift_with_direct_code(self):
        programmes = [{"id": "p1", "ten_khuyen_mai": "CT Một"},
                      {"id": "p2", "ten_khuyen_mai": "CT Hai"}]
        before = copy.deepcopy(programmes)
        source = bill([{"so_luong": 72, "promotion": programmes, "chiet_khau": 123},
                       {"is_km": True, "ctkmFull_id": "p1", "ctkmFull_ten_khuyen_mai": "CT Một", "so_luong": 12},
                       {"is_km": True, "ctkmFull_id": "p2", "ctkmFull_ten_khuyen_mai": "CT Hai", "so_luong": 24}])
        report, _, coverage = invoice_promotions(source, {}, OCT1)
        self.assertEqual(len(report), 3)
        self.assertEqual(report.iloc[0]["Mã CTKM"], "CT Hai; CT Một")
        self.assertEqual(list(report.iloc[1:]["Mã CTKM"]), ["CT Một", "CT Hai"])
        self.assertEqual(report.iloc[0]["Chiết khấu SP nguồn DMS"], 123)
        self.assertEqual(report.attrs["gift_rows"], 2)
        self.assertEqual(set(coverage["ID CT"]), {"p1", "p2"})
        self.assertEqual(list(coverage["Số đơn DMS"]), [1, 1])
        self.assertEqual(programmes, before)

    def test_invoice_uses_order_month_and_does_not_create_cash_for_gifts(self):
        source = bill([{"so_luong": 72, "promotion": [{"id": "p", "ten_khuyen_mai": "CT"}],
                        "ngay_dat": "2026-09-30 08:00:00"}])
        report, _, coverage = invoice_promotions(source, {}, OCT1)
        self.assertTrue(report.empty)
        self.assertTrue(coverage.empty)
        results, _ = calc.compute([qty_program()], calc.sold_lines(source, OCT1, OCT31), {})
        self.assertEqual(reward_ledger(results, OCT1)["Tiền mặt đủ điều kiện (đ)"].sum(), 0)

    def test_month_runner_includes_invoice_promotions_outside_bonus_purchase_skus(self):
        source = bill([{"ma_sp": "OTHER", "so_luong": 10, "promotion": [{"id": "p", "ten_khuyen_mai": "CT"}]}])
        with patch.object(bonus, "fetch_programs", return_value=[qty_program()]), \
                patch.object(bonus, "_bill_detail", return_value=source):
            manifest = {}
            frames = bonus._calc_month_frames(Mock(), bonus.load_config(), OCT1, OCT31, True,
                                             {}, "test", manifest, detail_config={})
        self.assertTrue(frames["BaoCao"].empty)
        self.assertEqual(len(frames["KhuyenMaiDonHang"]), 1)
        self.assertEqual(len(frames["ChuongTrinh"]), 2)
        self.assertEqual(manifest["reward_coverage"]["invoice_programs"], 1)
        self.assertEqual(manifest["reward_coverage"]["cash_program_levels"], 0)

    def test_export_keeps_cash_typed_and_long_program_names_readable(self):
        from openpyxl import load_workbook

        from promotion_detail import COLUMNS
        from promotion_workbook import write_detail_workbook

        source = bill([{"so_luong": 72, "thanh_tien": 400_000}])
        results, _ = calc.compute([cash_program()], calc.sold_lines(source, OCT1, OCT31), {})
        ledger = reward_ledger(results, OCT1)
        ledger["Tên CT"] = "Chương trình trả thưởng dành cho khách hàng " * 4
        with tempfile.TemporaryDirectory() as tmp:
            path = write_detail_workbook({"BaoCao": pd.DataFrame(columns=COLUMNS), "TraThuong": ledger},
                                         "report.xlsx", OCT1, Path(tmp))
            book = load_workbook(path)
            try:
                sheet = book["TraThuong"]
                headers = {c.value: c.column for c in sheet[1]}
                cash = sheet.cell(2, headers["Tiền mặt đủ điều kiện (đ)"])
                self.assertEqual((cash.value, cash.data_type, cash.number_format), (550_000, "n", "#,##0"))
                self.assertTrue(sheet.cell(2, headers["Tên CT"]).alignment.wrap_text)
                self.assertGreater(sheet.row_dimensions[2].height, 30)
                self.assertEqual(sheet.freeze_panes, "A2")
            finally:
                book.close()


if __name__ == "__main__":
    unittest.main()
